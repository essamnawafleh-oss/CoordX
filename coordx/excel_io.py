"""Read coordinate tables from Excel/CSV and write the reviewed result to Excel/CSV."""
from __future__ import annotations

import csv
import os
import re

from .model import Parcel, Point
from .textnorm import normalize, parse_float

_HEAD = {
    "no": ("no", "no.", "pt", "point", "pointno", "id", "name", "num", "number", "رقم", "رقمالنقطة", "النقطة", "نقطة", "م"),
    "x": ("x", "e", "east", "easting", "شرقي", "الشرقي", "سيني"),
    "y": ("y", "n", "north", "northing", "شمالي", "الشمالي", "صادي"),
    "z": ("z", "h", "elev", "elevation", "level", "منسوب", "المنسوب", "ارتفاع", "الارتفاع"),
}


def _head_role(v) -> str | None:
    if v is None:
        return None
    w = re.sub(r"[\s_:.\-()\[\]/]", "", normalize(str(v))).lower()
    if not w:
        return None
    for role, words in _HEAD.items():
        if w in words:
            return role
    for role, words in _HEAD.items():
        if any(len(k) > 2 and w.startswith(k) for k in words):
            return role
    return None


def _rows_to_parcels(rows: list[list], base: str, source: str) -> list[Parcel]:
    rows = [list(r) for r in rows]
    # Header: first row (within the first 10) that names at least X and Y
    roles = {}
    start = 0
    for i, r in enumerate(rows[:10]):
        rr = {}
        for ci, v in enumerate(r):
            role = _head_role(v)
            if role and role not in rr.values():
                rr[ci] = role
        if "x" in rr.values() and "y" in rr.values():
            roles, start = rr, i + 1
            break

    if not roles:
        # No header: guess from the numbers. A small integer column first -> point number.
        width = max((len(r) for r in rows), default=0)
        numeric = []
        for ci in range(width):
            vals = [parse_float(r[ci]) if ci < len(r) else None for r in rows]
            filled = [v for v in vals if v is not None]
            if filled and len(filled) >= 0.6 * sum(1 for r in rows if any(c not in (None, "") for c in r)):
                numeric.append((ci, filled))
        if numeric and all(float(v).is_integer() and abs(v) < 100000 for v in numeric[0][1]) and len(numeric) >= 3:
            roles[numeric[0][0]] = "no"
            numeric = numeric[1:]
        for (ci, _), role in zip(numeric, ("x", "y", "z")):
            roles[ci] = role
        if "x" not in roles.values():
            return []
        # skip leading text rows (titles)
        while start < len(rows) and all(parse_float(c) is None for c in rows[start]):
            start += 1

    col = {role: ci for ci, role in roles.items()}
    groups: list[list[Point]] = [[]]
    last_no = None
    for r in rows[start:]:
        get = lambda role: r[col[role]] if role in col and col[role] < len(r) else None
        x, y = parse_float(get("x")), parse_float(get("y"))
        if x is None and y is None:
            if groups[-1]:
                groups.append([])  # blank row separates parcels
            continue
        no_raw = get("no")
        no = "" if no_raw in (None, "") else str(no_raw).strip()
        if no.endswith(".0"):
            no = no[:-2]
        # numbering restarting at 1 starts a new parcel
        if no == "1" and last_no not in (None, "1") and groups[-1]:
            groups.append([])
        last_no = no
        groups[-1].append(Point(no=no, x=x, y=y, z=parse_float(get("z"))))

    groups = [g for g in groups if g]
    out = []
    for i, g in enumerate(groups, 1):
        for k, p in enumerate(g, 1):
            if not p.no:
                p.no = str(k)
        name = base if len(groups) == 1 else f"{base}-{i}"
        out.append(Parcel(name=name, points=g, source=source, method="excel"))
    return out


def read_table_file(path: str) -> list[Parcel]:
    ext = os.path.splitext(path)[1].lower()
    base = os.path.splitext(os.path.basename(path))[0]
    if ext in (".csv", ".txt"):
        with open(path, "rb") as f:
            raw = f.read()
        for enc in ("utf-8-sig", "cp1256", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        lines = [ln for ln in text.splitlines() if ln.strip()]
        head = "\n".join(lines[:20])
        # ';' and tab win over ',' because ',' may be the decimal separator
        delim = next((d for d in ("\t", ";", ",") if d in head), None)
        if delim:
            rows = list(csv.reader(lines, delimiter=delim))
        else:
            rows = [ln.split() for ln in lines]
        return _rows_to_parcels(rows, base, path)

    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    sheets = wb.worksheets
    for ws in sheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        name = base if len(sheets) == 1 else ws.title
        out.extend(_rows_to_parcels(rows, name, path))
    wb.close()
    return out


# ---------------------------------------------------------------- export

def _safe_sheet(name: str, used: set) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "_", name)[:28] or "Parcel"
    base, i = s, 2
    while s.lower() in used:
        s = f"{base[:25]}_{i}"
        i += 1
    used.add(s.lower())
    return s


def write_excel(parcels: list[Parcel], reports: list, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    bad_fill = PatternFill("solid", fgColor="F8D7DA")
    ok_font = Font(color="1E7B34")
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center")

    def header(ws, row, titles):
        for c, t in enumerate(titles, 1):
            cell = ws.cell(row=row, column=c, value=t)
            cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border

    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.rightToLeft = False
    header(summary, 1, ["Parcel / القطعة", "Points / النقاط", "Area m2 / المساحة",
                        "Perimeter m / المحيط", "Stated area / المساحة بالمخطط", "Issues / ملاحظات"])
    used = {"summary"}
    for r, (parcel, rep) in enumerate(zip(parcels, reports), 2):
        summary.append([parcel.name, len(parcel.valid_points()), round(rep.area, 2),
                        round(rep.perimeter, 2), parcel.stated_area, rep.issue_count])
        ws = wb.create_sheet(_safe_sheet(parcel.name, used))
        has_z = any(p.z is not None for p in parcel.points)
        cols = ["No / رقم", "X (E)", "Y (N)"] + (["Z"] if has_z else [])
        header(ws, 1, cols)
        for i, p in enumerate(parcel.points, 2):
            vals = [p.no, p.x, p.y] + ([p.z] if has_z else [])
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.border = border
                if c > 1:
                    cell.number_format = "0.00#"
                if (i - 2, ("no", "x", "y", "z")[c - 1]) in rep.cell_flags and rep.cell_flags[(i - 2, ("no", "x", "y", "z")[c - 1])][0] == "error":
                    cell.fill = bad_fill
        r0 = len(parcel.points) + 3
        header(ws, r0, ["Side / الضلع", "Length m / الطول", "On plan / بالمخطط", "Status / الحالة"])
        for k, s in enumerate(rep.sides, r0 + 1):
            ws.cell(row=k, column=1, value=s.name).border = border
            c = ws.cell(row=k, column=2, value=round(s.length, 2))
            c.border, c.number_format = border, "0.00"
            ws.cell(row=k, column=3, value=s.label).border = border
            st = ws.cell(row=k, column=4, value={"ok": "✓", "warn": "!", "none": ""}[s.status])
            st.border = border
            if s.status == "ok":
                st.font = ok_font
            elif s.status == "warn":
                st.fill = bad_fill
        k = r0 + len(rep.sides) + 2
        ws.cell(row=k, column=1, value="Area m2 / المساحة").font = Font(bold=True)
        ws.cell(row=k, column=2, value=round(rep.area, 2)).number_format = "0.00"
        ws.cell(row=k + 1, column=1, value="Perimeter m / المحيط").font = Font(bold=True)
        ws.cell(row=k + 1, column=2, value=round(rep.perimeter, 2)).number_format = "0.00"
        if parcel.stated_area:
            ws.cell(row=k + 2, column=1, value="Stated area / المساحة بالمخطط").font = Font(bold=True)
            ws.cell(row=k + 2, column=2, value=parcel.stated_area)
        for col, w in zip("ABCD", (16, 16, 16, 12)):
            ws.column_dimensions[col].width = w
    for col, w in zip("ABCDEF", (22, 14, 18, 18, 24, 16)):
        summary.column_dimensions[col].width = w
    wb.save(path)


def write_csv(parcels: list[Parcel], path: str):
    """PNEZD comma-delimited, as Civil 3D's point import expects."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for parcel in parcels:
            for p in parcel.valid_points():
                w.writerow([p.no, f"{p.y:.3f}", f"{p.x:.3f}", "" if p.z is None else f"{p.z:.3f}", parcel.name])
